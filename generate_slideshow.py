from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SUPPORTED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff", ".webp"}


def collect_images(photos_dir: Path) -> list[Path]:
    return sorted(
        path for path in photos_dir.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def create_title_slide(presentation: Presentation, title: str, subtitle: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def fit_image(image_path: Path, slide_width: int, slide_height: int, margin_emu: int) -> tuple[int, int]:
    with Image.open(image_path) as image:
        width, height = image.size

    max_width = slide_width - (margin_emu * 2)
    max_height = slide_height - (margin_emu * 2)
    scale = min(max_width / width, max_height / height)
    return int(width * scale), int(height * scale)


def create_photo_slide(presentation: Presentation, image_path: Path) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_background(slide, RGBColor(13, 27, 42))

    margin_emu = int(Inches(0.35))
    image_width, image_height = fit_image(
        image_path, presentation.slide_width, presentation.slide_height, margin_emu
    )
    left = int((presentation.slide_width - image_width) / 2)
    top = int((presentation.slide_height - image_height) / 2)

    frame = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        left - Inches(0.08),
        top - Inches(0.08),
        image_width + Inches(0.16),
        image_height + Inches(0.16),
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.fill.background()

    slide.shapes.add_picture(str(image_path), left, top, width=image_width, height=image_height)


def create_closing_slide(presentation: Presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_background(slide, RGBColor(13, 27, 42))

    left = int(Inches(1))
    top = int(Inches(2.2))
    width = presentation.slide_width - (left * 2)
    textbox = slide.shapes.add_textbox(left, top, width, int(Inches(2)))
    text_frame = textbox.text_frame
    paragraph = text_frame.paragraphs[0]
    paragraph.text = "Thank you for celebrating the year with us!"
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(255, 255, 255)


def build_presentation(photos_dir: Path, output_path: Path, title: str, subtitle: str) -> int:
    images = collect_images(photos_dir)
    if not images:
        raise ValueError(f"No supported images were found in {photos_dir}.")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    create_title_slide(presentation, title, subtitle)
    for image_path in images:
        create_photo_slide(presentation, image_path)
    create_closing_slide(presentation)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return len(images)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a presentable end-of-year PowerPoint from a folder of photos."
    )
    parser.add_argument("--photos-dir", default="photos", help="Folder containing the photos to include.")
    parser.add_argument(
        "--output",
        default="output/end-of-year-celebration.pptx",
        help="Where to save the generated PowerPoint file.",
    )
    parser.add_argument("--title", default="End of Year Celebration", help="Title slide heading.")
    parser.add_argument(
        "--subtitle",
        default="A slideshow of favorite memories from the school year",
        help="Subtitle shown on the title slide.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    photos_dir = Path(args.photos_dir)
    try:
        if not photos_dir.exists():
            raise FileNotFoundError(f"Photo folder not found: {photos_dir}")

        output_path = Path(args.output)
        image_count = build_presentation(photos_dir, output_path, args.title, args.subtitle)
    except (FileNotFoundError, ValueError) as error:
        raise SystemExit(str(error)) from error

    print(f"Created {output_path} with {image_count} photo slides.")


if __name__ == "__main__":
    main()
